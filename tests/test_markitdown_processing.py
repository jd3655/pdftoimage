import importlib.util
import json
import zipfile
from pathlib import Path

import pytest

markitdown_spec = importlib.util.find_spec("markitdown")
docx_spec = importlib.util.find_spec("docx")
pptx_spec = importlib.util.find_spec("pptx")
openpyxl_spec = importlib.util.find_spec("openpyxl")

if not (markitdown_spec and docx_spec and pptx_spec and openpyxl_spec):  # pragma: no cover - environment dependent
    pytest.skip("Skipping MarkItDown integration tests; dependencies not available", allow_module_level=True)

from markitdown_processing import MarkItDownOptions, convert_batch_to_outputs


def _create_docx(path: Path):
    from docx import Document

    doc = Document()
    doc.add_heading("Sample Document", level=1)
    doc.add_paragraph("Hello from python-docx")
    doc.save(path)


def _create_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    textbox = slide.shapes.add_textbox(left=0, top=0, width=prs.slide_width, height=prs.slide_height // 4)
    textbox.text = "Hello from python-pptx"
    prs.save(path)


def _create_xlsx(path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet["A1"] = "Hello"
    sheet["B1"] = "MarkItDown"
    sheet.append(["Another", "Row"])
    wb.save(path)


def _create_textual_files(tmp_path: Path):
    html = tmp_path / "page.html"
    html.write_text("<html><body><p>HTML content</p></body></html>", encoding="utf-8")
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("col1,col2\n1,2\n", encoding="utf-8")
    json_file = tmp_path / "data.json"
    json_file.write_text(json.dumps({"key": "value"}), encoding="utf-8")
    xml_file = tmp_path / "data.xml"
    xml_file.write_text("<root><item>value</item></root>", encoding="utf-8")
    return [html, csv_file, json_file, xml_file]


def test_markitdown_batch_conversion(tmp_path: Path):
    docx_path = tmp_path / "sample.docx"
    pptx_path = tmp_path / "slides.pptx"
    xlsx_path = tmp_path / "sheet.xlsx"
    _create_docx(docx_path)
    _create_pptx(pptx_path)
    _create_xlsx(xlsx_path)

    textual_files = _create_textual_files(tmp_path)
    archive_path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(archive_path, "w") as zf:
        for file_path in textual_files:
            zf.write(file_path, file_path.name)

    options = MarkItDownOptions()
    output_zip, manifest_info, _ = convert_batch_to_outputs(
        [str(docx_path), str(pptx_path), str(xlsx_path)],
        str(archive_path),
        options,
    )

    assert output_zip.exists()
    with zipfile.ZipFile(output_zip, "r") as zf:
        manifest = json.loads(zf.read("manifest.json"))
        markdown_manifest = json.loads(zf.read("markdown_manifest.json"))
        markdown_files = [name for name in zf.namelist() if name.endswith(".md")]

    expected_entries = 3 + len(textual_files)
    assert len(manifest) == expected_entries
    assert len(markdown_manifest) >= expected_entries
    assert all(entry["status"] == "success" for entry in manifest)
    assert markdown_files, "No markdown outputs were packaged."

    # Ensure markdown content is non-empty for at least one file.
    with zipfile.ZipFile(output_zip, "r") as zf:
        sample_md = zf.read(markdown_files[0]).decode("utf-8")
    assert sample_md.strip()
